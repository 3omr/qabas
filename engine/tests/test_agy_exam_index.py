"""Source-grounded agy question extraction and resumable index batches."""

import hashlib
import io
import json
import os
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

SCRIPTS_DIR = Path(__file__).resolve().parent.parent / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

import agy_exam_index  # noqa: E402
from exam_index import carry_over_repairs, write_index  # noqa: E402
from exam_question_review import review_paper as REVIEW_PAPER  # noqa: E402
from lecture_registry import list_module_files  # noqa: E402
from module_registry import load_module  # noqa: E402
from question_provenance import paper_backed_index  # noqa: E402


class AgyExamIndexTests(unittest.TestCase):
    def setUp(self):
        backoff = patch("exam_index_recovery.wait_backoff")
        self.backoff = backoff.start()
        self.addCleanup(backoff.stop)
        review = patch("exam_question_review.review_paper", side_effect=lambda _snapshot, questions, *_args, **_options: (questions, "done", None))
        review.start()
        self.addCleanup(review.stop)
        self.temporary = tempfile.TemporaryDirectory(prefix="qabas-agy-exams-")
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name) / "modules" / "pediatric"
        self.questions = self.root / "Questions"
        self.questions.mkdir(parents=True)
        (self.root / "Lecture").mkdir()
        (self.root / "Transcripts").mkdir()
        (self.root / "module.json").write_text(json.dumps({
            "schema_version": 1, "module_id": "pediatric", "display_name": "Pediatric",
            "notebook": {"id": "test-notebook"},
        }), encoding="utf-8")
        self.module = load_module(self.root)

    def test_nested_same_name_papers_keep_independent_identity_through_publication_and_edits(self):
        import mcp_server
        from desktop_library import question_index_status
        from exam_preparation import exam_source_files
        from exam_year_overrides import load_year_overrides, set_exam_year
        from lecture_registry import remove_file, rename_file
        from library_trash import restore_trash
        from question_provenance import assessment_catalog, assessment_verified_years
        from run_transcription import assessment_entry, build_exam_index

        names = ["ENDS/Final 25.txt", "Finals/Final 25.txt", "Bank.txt", "2024/Undated.txt", "Questions/Extra.txt", "notes.json.txt"]
        for name in [*names, ".hidden/ignored.txt", "ENDS/.ignored.txt"]:
            path = self.questions / name
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text("1. Explain dehydration.\n", encoding="utf-8")
        (self.questions / "notes.json").write_text("{}")
        files = exam_source_files(self.module)
        self.assertEqual({path.relative_to(self.questions).as_posix() for path in files}, set(names))
        with patch("agy_exam_index.agy_writer.request_json", return_value=self._written_response(files[0])):
            index, target = build_exam_index(self.module, {"exam_index_workers": 1})
        self.assertEqual({source["file"] for source in index["sources"]}, set(names))
        self.assertEqual(index["failed_sources"], [])
        verified = paper_backed_index(index, {}, "pediatric", self.questions)
        self.assertEqual({o["source"] for q in verified["questions"].values() for o in q["occurrences"]}, set(names))
        self.assertFalse(mcp_server._exam_index_is_stale(index, target, self.module, require_complete=True))
        self.assertEqual(question_index_status(self.module)["exam_index"], "built")
        with patch("remote_inventory.module_inventory", return_value=SimpleNamespace(
            sources=[], warning=None, remote_as_of=None, available=False,
        )):
            inventory = list_module_files(self.module)["files"]
        self.assertEqual({item["name"] for item in inventory}, set(names))
        self.assertTrue(all(item["indexed"] for item in inventory))
        self.assertEqual(next(item for item in inventory if item["name"] == "2024/Undated.txt")["year"], None)
        # The MCP registry resolves the same fixture module through its normal workspace id.
        page = json.loads(mcp_server._list_exam_questions({"module": "pediatric", "path": "Questions/ENDS/Final 25.txt"}, self.root.parent.parent))
        self.assertEqual(page["total"], 1)
        found = json.loads(mcp_server._find_questions({"module": "pediatric", "lecture": "Dehydration"}, self.root.parent.parent))
        self.assertEqual({item["path"] for entry in found["entries"] for item in entry["source_papers"]}, {"Questions/" + name for name in names})
        from phase_validation import _catalog_matches
        both = assessment_catalog(self.root, {"assessment_sources": [assessment_entry(self.questions / name) for name in names[:2]]})
        self.assertEqual([entry["canonical_name"] for entry in _catalog_matches("Questions/ENDS/Final 25.txt", both)], [names[0]])
        self.assertEqual(_catalog_matches("Final 25.txt", both), [])
        from provenance_audit import supported_years
        self.assertEqual(supported_years("Explain dehydration.", {"2024/Undated.txt": "1. Explain dehydration."}), ())
        set_exam_year(self.module, "Questions/ENDS/Final 25.txt", 2022)
        catalog = assessment_catalog(self.root, {"assessment_sources": [assessment_entry(self.questions / names[0])]})
        self.assertEqual(assessment_verified_years(catalog), {2022})
        rename_file(self.module, "Questions/ENDS/Final 25.txt", "Renamed.txt")
        self.assertEqual(load_year_overrides(self.questions), {"ENDS/Renamed.txt": 2022})
        self.assertTrue((self.questions / names[1]).exists())
        removed = remove_file(self.module, "Questions/ENDS/Renamed.txt")
        self.assertFalse((self.questions / "ENDS/Renamed.txt").exists())
        restore_trash(self.module, removed["id"])
        self.assertTrue((self.questions / "ENDS/Renamed.txt").exists())
        self.assertTrue(mcp_server._exam_index_is_stale(index, target, self.module, require_complete=True))

    def test_spreadsheet_derived_sibling_is_preserved_and_indexed_only_as_original(self):
        from exam_preparation import exam_source_files
        from openpyxl import Workbook
        from run_transcription import build_exam_index

        source = self.questions / "Banks/Pediatrics.xlsx"
        source.parent.mkdir()
        workbook = Workbook()
        workbook.active.append(["Question"])
        workbook.active.append(["Explain dehydration."])
        workbook.save(source)
        workbook.close()
        original = source.read_bytes()
        stale = source.with_name(source.name + ".txt")
        stale.write_text("old student bytes", encoding="utf-8")
        self.assertEqual(exam_source_files(self.module), [source])
        with patch("agy_exam_index.agy_writer.request_json", return_value=self._written_response(source)):
            index, _ = build_exam_index(self.module, {})
        self.assertEqual(index["failed_sources"], [])
        self.assertEqual([item["file"] for item in index["sources"]], ["Banks/Pediatrics.xlsx"])
        self.assertEqual(index["sources"][0]["questions"], 1)
        self.assertEqual(source.read_bytes(), original)
        receipt = json.loads((self.root / ".transcriber-cache/exam-texts.json").read_text())["Banks/Pediatrics.xlsx"]
        self.assertEqual((self.root / receipt["preserved_text"]).read_text(), "old student bytes")
        self.assertIn("Explain dehydration.", stale.read_text())
        self.assertEqual(exam_source_files(self.module), [source])
        with patch("agy_exam_index.agy_writer.request_json", side_effect=AssertionError("cache should be reused")):
            cached, _ = build_exam_index(self.module, {})
        self.assertEqual(cached["sources"], index["sources"])

    def test_large_numbered_paper_batches_before_requests_and_reuses_only_its_layout(self):
        from threading import Barrier

        source = self.questions / "Khalsa.txt"
        source.write_text("\n".join(f"{n}) Which poison in case {n}? A. Opium. B. Lead." for n in range(1, 75)), encoding="utf-8")
        batches = []
        concurrent = Barrier(3, timeout=5)

        def writer(prompt, _schema, **options):
            self.assertEqual(options["timeout"], 137)
            units = json.loads(prompt.split("Evidence units: ", 1)[1])
            core = [unit for unit in units if unit["coverage"] == "required"]
            ids = [unit["id"] for unit in core]
            batches.append(ids)
            if len(core) == 20:
                concurrent.wait()
            return {"covered_unit_ids": ids, "questions": [
                self._written_response(source, unit_ids=[unit["id"]], kind="mcq",
                    options=[{"label": "A", "text": "Opium."}, {"label": "B", "text": "Lead."}],
                    stem=unit["text"].split(") ", 1)[1].split(" A.", 1)[0])["questions"][0] for unit in core]}

        config = {"exam_index_request_timeout_seconds": 137}
        with patch("agy_exam_index.agy_writer.request_json", side_effect=writer) as request:
            index = agy_exam_index.build_index(self.module, [source], config=config)
            self.assertEqual(index["failed_sources"], [])
            self.assertEqual(index["sources"][0]["questions"], 74)
            self.assertEqual(sorted(map(len, batches)), [14, 20, 20, 20])
            self.assertEqual(request.call_count, 4)
            agy_exam_index.build_index(self.module, [source], config=config)
            self.assertEqual(request.call_count, 4)
            # The number of batches stays four, but each batch covers different units.
            agy_exam_index.build_index(self.module, [source], config={**config, "exam_index_batch_questions": 19})
            self.assertEqual(request.call_count, 8)
            state_path = self.questions / agy_exam_index.STATE_NAME
            state = json.loads(state_path.read_text())
            del state["files"][source.name]["batch_layout"]
            state_path.write_text(json.dumps(state))
            agy_exam_index.build_index(self.module, [source], config={**config, "exam_index_batch_questions": 19})
            self.assertEqual(request.call_count, 12)

    def test_question_start_formats_and_unnumbered_char_batches(self):
        for prefix in ("1)", "1.", "1-", "Q1", "Q 1:", "١)", "۲.", "1)Cl"):
            with self.subTest(prefix=prefix):
                units = agy_exam_index._unit_text("\n".join(f"{prefix} Explain poisoning." for _ in range(74)))
                batches = agy_exam_index._unit_batches(units, 12000)
                self.assertEqual([len(core) for core, _ in batches], [20, 20, 20, 14])
                self.assertEqual(batches[1][1][:5], units[15:20])
                self.assertEqual(batches[1][1][-5:], units[40:45])
        units = agy_exam_index._unit_text("\n".join("1.5 mg is unnumbered question text." for _ in range(74)))
        self.assertEqual(len(agy_exam_index._unit_batches(units, 12000)), 1)
        self.assertEqual([len(core) for core, _ in agy_exam_index._unit_batches(units, len(units[0].text) * 2)], [2] * 37)
        rows = [agy_exam_index.SourceUnit(f"row{n}", {"type": "spreadsheet_row", "row": n}, "Question")
                for n in range(1, 76)]
        self.assertEqual([sum(agy_exam_index._question_starts(unit) for unit in core)
                          for core, _ in agy_exam_index._unit_batches(rows, 12000)], [20, 20, 20, 14])

    def test_index_request_budgets_reject_invalid_config(self):
        for key in ("exam_index_batch_questions", "exam_index_request_timeout_seconds"):
            for value in (0, -1, True, "20", 1.5, None):
                with self.subTest(key=key, value=value), self.assertRaisesRegex(ValueError, key):
                    agy_exam_index.index_settings({key: value})

    def test_partial_inventory_keeps_matching_papers_and_failures_independent(self):
        from desktop_library import question_index_status
        from exam_index import SCHEMA_VERSION
        from mcp_server import _exam_index_is_stale

        papers = [self.questions / name for name in ("A.txt", "B.txt", "C.txt")]
        for paper in papers:
            paper.write_text("1. Which symptom?\nA. Fever\nB. Pain\n", encoding="utf-8")
        bank = {"schema_version": SCHEMA_VERSION, "extractor": "agy", "module": "pediatric",
                "sources": [{"file": "A.txt", "sha256": hashlib.sha256(papers[0].read_bytes()).hexdigest(),
                             "questions": 3, "needs_review": 1}], "questions": {},
                "failed_sources": [{"file": "B.txt", "sha256": hashlib.sha256(papers[1].read_bytes()).hexdigest(),
                                    "error": "agy returned no grounded questions", "cause": "model"}]}
        write_index(bank, self.questions)
        with patch("remote_inventory.module_inventory", return_value=SimpleNamespace(
            sources=[], warning=None, remote_as_of=None, available=False,
        )):
            inventory = {item["name"]: item for item in list_module_files(self.module)["files"]}
        self.assertTrue(inventory["A.txt"]["indexed"])
        self.assertEqual(inventory["A.txt"]["question_count"], 3)
        self.assertEqual(inventory["A.txt"]["question_review_count"], 1)
        self.assertTrue(inventory["B.txt"]["index_failed"])
        self.assertEqual(inventory["B.txt"]["index_error"], "agy returned no grounded questions")
        self.assertEqual(inventory["B.txt"]["index_failure_cause"], "model")
        self.assertFalse(inventory["C.txt"]["indexed"])
        self.assertFalse(inventory["C.txt"]["index_failed"])
        self.assertEqual(question_index_status(self.module), {
            "exam_index": "partial", "question_files": 3,
            "question_index_counts": {"indexed": 1, "failed": 1, "pending": 1},
        })
        self.assertFalse(_exam_index_is_stale(bank, self.questions / "exam-index.json", self.module, papers[0]))
        papers[0].write_text("changed paper", encoding="utf-8")
        statuses = {}
        summary = question_index_status(self.module, statuses)
        self.assertFalse(statuses[papers[0]]["indexed"])
        self.assertNotIn("question_count", statuses[papers[0]])
        self.assertEqual(summary["exam_index"], "stale")
        self.assertEqual(summary["question_index_counts"], {"indexed": 0, "failed": 1, "pending": 2})
        self.assertTrue(_exam_index_is_stale(bank, self.questions / "exam-index.json", self.module, papers[0]))
        papers[1].write_text("changed failed paper", encoding="utf-8")
        question_index_status(self.module, statuses)
        self.assertFalse(statuses[papers[1]]["index_failed"])
        self.assertNotIn("index_error", statuses[papers[1]])
        self.assertNotIn("index_failure_cause", statuses[papers[1]])

    def test_detailed_failures_keep_inventory_compatible_with_existing_consumers(self):
        from exam_index import SCHEMA_VERSION

        source = self.questions / "paper.txt"
        source.write_text("Explain dehydration.")
        for cause, presentation in (("auth", "auth"), ("quota", "model"), ("source_unreadable", "source"),
                                    ("source_ocr_unavailable", "source"), ("model_unavailable", "model"),
                                    ("timeout", "model"), ("invalid_output", "model")):
            with self.subTest(cause=cause):
                bank = {"schema_version": SCHEMA_VERSION, "extractor": "agy", "module": self.module.module_id,
                        "sources": [], "questions": {}, "failed_sources": [agy_exam_index.failed_source(source, "Plain last error", cause, 3)]}
                write_index(bank, self.questions)
                with patch("remote_inventory.module_inventory", return_value=SimpleNamespace(
                    sources=[], warning=None, remote_as_of=None, available=False,
                )):
                    paper = list_module_files(self.module)["files"][0]
                self.assertEqual(paper["index_failure_cause"], presentation)
                self.assertEqual(paper["index_error"], "Plain last error")
                self.assertEqual(bank["failed_sources"][0]["cause"], cause)
                self.assertEqual(bank["failed_sources"][0]["attempts"], 3)

    def test_kept_older_failure_without_a_cause_gets_stable_metadata(self):
        kept = self.questions / "kept.txt"
        kept.write_text("Retained failed paper.")
        source = self.questions / "ready.txt"
        source.write_text("Explain dehydration.")
        with patch("agy_exam_index.agy_writer.request_json", return_value=self._written_response(source)):
            index = agy_exam_index.build_index(self.module, [source], config={},
                failed_sources=[{"file": kept.name, "sha256": agy_exam_index._sha256(kept), "error": "Old diagnostic"}],
                report_progress=lambda _done, _total, _message, **_details: None)
        self.assertEqual(index["failed_sources"][0]["cause"], "invalid_output")
        self.assertEqual(index["failed_sources"][0]["attempts"], 0)
        self.assertEqual(index["failed_sources"][0]["error"], "Old diagnostic")

    def test_legacy_bank_remains_readable_with_a_new_unprepared_paper(self):
        from desktop_library import question_index_status
        from exam_index import build_index
        from mcp_server import _find_questions, _list_exam_questions, _questions_status

        source = self.questions / "2023.txt"
        source.write_text("1. Which symptom occurs in pneumonia?\nA. Fever\nB. Pain\n", encoding="utf-8")
        write_index(build_index(self.questions, "pediatric"), self.questions)
        (self.questions / "new.pdf").write_bytes(b"unread scan")
        summary = question_index_status(self.module)
        self.assertEqual(summary["exam_index"], "legacy")
        self.assertEqual(summary["question_index_counts"], {"indexed": 0, "failed": 0, "pending": 2})
        statuses = {}
        question_index_status(self.module, statuses)
        self.assertTrue(statuses[source]["index_legacy"])
        self.assertFalse(statuses[source]["indexed"])
        self.assertEqual(_questions_status(self.module), "indexed")
        lecture = json.loads(_find_questions({"module": "pediatric", "lecture": "pneumonia"}, self.root.parent.parent))
        self.assertEqual(lecture["total"], 1)
        file_questions = json.loads(_list_exam_questions({"module": "pediatric", "path": "Questions/2023.txt"}, self.root.parent.parent))
        self.assertEqual(file_questions["total"], 1)
        self.assertIn("pneumonia", file_questions["questions"][0]["stem"])

    def test_capacity_failures_split_recursively_and_reuse_cached_leaves(self):
        source = self.questions / "dense.txt"
        source.write_text("\n".join(f"Explain printed finding number {n}." for n in range(8)), encoding="utf-8")
        failures = [
            agy_exam_index.agy_writer.AgyWriterError("agy timed out after 600s."),
            agy_exam_index.agy_writer.AgyWriterError("MAX_TOKENS output limit"),
            agy_exam_index.agy_writer.AgyWriterError("response truncated"),
            agy_exam_index.agy_writer.AgyWriterError("agy returned invalid JSON: unexpected end"),
            agy_exam_index.agy_writer.AgyProposalError("invalid proposal JSON", "{broken"),
        ]
        for failure in failures:
            with self.subTest(failure=str(failure)):
                (self.questions / agy_exam_index.STATE_NAME).unlink(missing_ok=True)
                seen = []
                progress = []
                def fake(prompt, _schema, seen=seen, failure=failure, **_options):
                    units, _end = json.JSONDecoder().raw_decode(prompt.split("Evidence units: ", 1)[1])
                    core = [unit for unit in units if unit["coverage"] == "required"]
                    seen.append(tuple(unit["id"] for unit in core))
                    if len(core) > 2:
                        raise failure
                    return {"covered_unit_ids": [unit["id"] for unit in core], "questions": [
                        self._written_response(source, unit_ids=[unit["id"]], stem=unit["text"])["questions"][0]
                        for unit in core
                    ]}
                with patch("agy_exam_index.agy_writer.request_json", side_effect=fake):
                    index = agy_exam_index.build_index(self.module, [source], config={"exam_index_workers": 3, "exam_index_timeout_limit": 10},
                        report_progress=lambda _done, _total, message, messages=progress: messages.append(message))
                self.assertEqual(index["failed_sources"], [])
                self.assertEqual([question["stem"] for question in index["questions"].values()], source.read_text().splitlines())
                self.assertIn(tuple(f"L{n:05d}" for n in range(1, 9)), seen)
                self.assertTrue(any("part 1 of 1 · split 1/2" in message for message in progress))
                state = json.loads((self.questions / agy_exam_index.STATE_NAME).read_text())
                batches = state["files"][source.name]["batches"]
                self.assertEqual({key for key, value in batches.items() if value.get("split")}, {"0", "0/1", "0/2"})
                self.assertEqual({key for key, value in batches.items() if "payload" in value},
                                 {"0/1/1", "0/1/2", "0/2/1", "0/2/2"})
                with patch("agy_exam_index.agy_writer.request_json", side_effect=AssertionError("cached split was requested")):
                    resumed = agy_exam_index.build_index(self.module, [source], config={"exam_index_workers": 1})
                self.assertEqual(index, resumed)

    def test_single_unit_timeout_isolated_to_its_paper_and_other_leaves_survive(self):
        first, second = self.questions / "dense.txt", self.questions / "other.txt"
        first.write_text("\n".join(f"Source note {n}." for n in range(4)), encoding="utf-8")
        second.write_text("Other paper note.", encoding="utf-8")
        import threading

        seen = []
        answered = set()
        siblings_done = threading.Condition()
        def fake(prompt, _schema, **_options):
            units = json.loads(prompt.split("Evidence units: ", 1)[1])
            core = [unit["id"] for unit in units if unit["coverage"] == "required"]
            if "Source file: dense.txt" in prompt:
                seen.append(core)
                if core == ["L00002"]:
                    # A failed paper stops its running batches; let the other leaves finish first so
                    # the cache this test resumes from does not depend on thread timing.
                    with siblings_done:
                        siblings_done.wait_for(lambda: {"L00001", "L00003", "L00004"} <= answered, timeout=5)
                if len(core) > 1 or core == ["L00002"]:
                    raise agy_exam_index.agy_writer.AgyWriterError("agy timed out after 600s.")
                with siblings_done:
                    answered.update(core)
                    siblings_done.notify_all()
            return {"covered_unit_ids": core, "questions": []}
        config = {"exam_index_timeout_limit": 10}
        with patch("agy_exam_index.agy_writer.request_json", side_effect=fake):
            index = agy_exam_index.build_index(self.module, [first, second], config=config)
            seen.clear()
            resumed = agy_exam_index.build_index(self.module, [first, second], config=config)
        self.assertEqual(index["sources"], resumed["sources"])
        self.assertGreater(resumed["failed_sources"][0]["attempts"], index["failed_sources"][0]["attempts"])
        self.assertEqual([source["file"] for source in index["sources"]], ["other.txt"])
        self.assertEqual(index["failed_sources"][0]["file"], "dense.txt")
        self.assertIn("timed out", index["failed_sources"][0]["error"])
        self.assertEqual(seen, [["L00002"], ["L00002"]])

    def test_timeout_limit_fails_the_paper_instead_of_splitting_again(self):
        source = self.questions / "dense.txt"
        source.write_text("\n".join(f"Source note {n}." for n in range(8)), encoding="utf-8")
        timeout = agy_exam_index.agy_writer.AgyWriterError("agy timed out after 600s.")
        for limit, calls in ((None, 7), (1, 2)):
            with self.subTest(limit=limit):
                (self.questions / agy_exam_index.STATE_NAME).unlink(missing_ok=True)
                events = []
                config = {} if limit is None else {"exam_index_timeout_limit": limit}
                with patch("agy_exam_index.agy_writer.request_json", side_effect=timeout) as request:
                    index = agy_exam_index.build_index(self.module, [source], config=config,
                        report_progress=lambda done, total, message, events=events, **details: events.append(details))
                if limit == 1:
                    self.assertEqual(request.call_count, calls)
                else:
                    self.assertGreaterEqual(request.call_count, 4)
                    self.assertLessEqual(request.call_count, calls)
                [failure] = index["failed_sources"]
                self.assertEqual(failure["cause"], "timeout")
                self.assertEqual(failure["attempts"], request.call_count)
                self.assertIn(f"agy stopped responding ({limit or 2} timeouts)", failure["error"])
                self.assertEqual([event["cause"] for event in events if event.get("phase") == "failed"], ["timeout"])
        for invalid in (0, -1, True, 1.5, "2"):
            with self.assertRaisesRegex(ValueError, "exam_index_timeout_limit"):
                agy_exam_index.timeout_limit({"exam_index_timeout_limit": invalid})

    def test_signed_out_agy_fails_remaining_papers_without_new_requests(self):
        cached, first, second = (self.questions / name for name in ("Cached.txt", "First.txt", "Second.txt"))
        for source in (cached, first, second):
            source.write_text("Source note one.\n", encoding="utf-8")
        def covered(prompt, _schema, **_options):
            units = json.loads(prompt.split("Evidence units: ", 1)[1])
            return {"covered_unit_ids": [unit["id"] for unit in units if unit["coverage"] == "required"], "questions": []}
        config = {"exam_index_workers": 1}
        with patch("agy_exam_index.agy_writer.request_json", side_effect=covered):
            agy_exam_index.build_index(self.module, [cached], config=config)
        events = []
        signed_out = agy_exam_index.agy_writer.AgyWriterError("agy failed: You are not signed in. Sign in with Google.")
        with patch("agy_exam_index.agy_writer.request_json", side_effect=signed_out) as request:
            index = agy_exam_index.build_index(self.module, [cached, first, second], config=config,
                report_progress=lambda done, total, message, **details: events.append(details))
        self.assertEqual(request.call_count, 1)
        self.assertEqual([source["file"] for source in index["sources"]], ["Cached.txt"])
        causes = {failure["file"]: failure["cause"] for failure in index["failed_sources"]}
        self.assertEqual(causes, {"First.txt": "auth", "Second.txt": "auth"})
        self.assertEqual({event["file"]: event["cause"] for event in events if event.get("phase") == "failed"}, causes)
        self.assertEqual({event["file"] for event in events if event.get("phase") in {"done", "failed"}},
                         {"Cached.txt", "First.txt", "Second.txt"})

    def test_issue_61_account_quota_retains_cached_papers_and_blocks_queued_requests(self):
        cached, first, second = (self.questions / name for name in ("Cached.txt", "First.txt", "Second.txt"))
        for source in (cached, first, second):
            source.write_text("1. Explain dehydration.\n", encoding="utf-8")
        config = {"exam_index_workers": 1}
        with patch("agy_exam_index.agy_writer.request_json", return_value=self._written_response(cached)):
            completed = agy_exam_index.build_index(self.module, [cached], config=config)
        diagnostic = "RESOURCE_EXHAUSTED (code 429): Individual quota reached. Resets in 167h37m19s."
        with patch("agy_exam_index.agy_writer.request_json", side_effect=agy_exam_index.agy_writer.AgyWriterError(diagnostic)) as request:
            index = agy_exam_index.build_index(self.module, [cached, first, second], config=config)
        self.assertEqual(request.call_count, 1)
        self.assertEqual(index["questions"], completed["questions"])
        self.assertEqual([source["file"] for source in index["sources"]], [cached.name])
        self.assertEqual({source["file"] for source in index["failed_sources"]}, {first.name, second.name})
        self.assertEqual(index["interruption"]["kind"], "quota")
        self.assertEqual(index["interruption"]["raw_diagnostic"], diagnostic)
        self.assertIsNotNone(index["interruption"]["reset_at"])
        with patch("agy_exam_index.agy_writer.request_json", return_value=self._written_response(first)) as request:
            resumed = agy_exam_index.build_index(self.module, [cached, first, second], config=config)
        self.assertNotIn("interruption", resumed)
        self.assertEqual(resumed["failed_sources"], [])

    def test_failed_paper_cancels_its_running_batches(self):
        import threading
        import time

        import cancellation

        source = self.questions / "dense.txt"
        source.write_text("\n".join(f"Source note {n}." for n in range(4)), encoding="utf-8")
        started = threading.Event()
        def fake(prompt, _schema, **_options):
            units = json.loads(prompt.split("Evidence units: ", 1)[1])
            core = [unit["id"] for unit in units if unit["coverage"] == "required"]
            if core == ["L00001"]:
                started.wait(5)
                raise agy_exam_index.agy_writer.AgyWriterError("agy failed: unsupported request")
            started.set()
            cancellation.run([sys.executable, "-c", "import time; time.sleep(30)"], timeout=60)
            return {"covered_unit_ids": core, "questions": []}
        began = time.monotonic()
        with patch("agy_exam_index.agy_writer.request_json", side_effect=fake):
            index = agy_exam_index.build_index(self.module, [source], config={"exam_index_batch_chars": 18, "exam_index_workers": 2})
        self.assertLess(time.monotonic() - began, 10)
        self.assertEqual([failure["file"] for failure in index["failed_sources"]], ["dense.txt"])

    def test_split_overlap_retains_neighbours_and_single_unit_core(self):
        units = [agy_exam_index.SourceUnit(f"L{n:05d}", {"type": "line", "line": n}, str(n)) for n in range(20)]
        task = agy_exam_index.BatchTask(2, (), units[5:15], units)
        left, right = task.children()
        self.assertEqual(left.units, units[:15])
        self.assertEqual(right.units, units[5:20])
        self.assertEqual([left.key, right.key], ["2/1", "2/2"])
        while len(left.core) > 1:
            left = left.children()[0]
        self.assertEqual(left.core, units[5:6])
        self.assertEqual(left.units, units[:11])

    def test_authentication_errors_do_not_split_large_batches(self):
        source = self.questions / "dense.txt"
        source.write_text("Source note one.\nSource note two.\n", encoding="utf-8")
        with patch("agy_exam_index.agy_writer.request_json", side_effect=agy_exam_index.agy_writer.AgyWriterError("authentication required")) as request:
            index = agy_exam_index.build_index(self.module, [source], config={})
        self.assertEqual(request.call_count, 1)
        self.assertEqual(len(index["failed_sources"]), 1)

    def test_failures_name_whether_the_paper_or_the_model_run_failed(self):
        launch = self.questions / "End_2021.txt"
        launch.write_text("Q1. Source question?\nA. one\nB. two\n", encoding="utf-8")
        empty = self.questions / "Blank.txt"
        empty.write_text("\n", encoding="utf-8")
        events = []
        error = agy_exam_index.agy_writer.AgyWriterError("Could not run agy: [WinError 206] The filename or extension is too long")
        with patch("agy_exam_index.agy_writer.request_json", side_effect=error):
            index = agy_exam_index.build_index(self.module, [empty, launch], config={},
                report_progress=lambda done, total, message, **details: events.append(details))
        causes = {failure["file"]: failure["cause"] for failure in index["failed_sources"]}
        self.assertEqual(causes, {"Blank.txt": "source_unreadable", "End_2021.txt": "invalid_output"})
        self.assertEqual({event["file"]: event["cause"] for event in events if event.get("phase") == "failed"}, causes)

    def test_configured_source_budget_and_invalid_settings(self):
        from math import ceil

        units = [agy_exam_index.SourceUnit(f"L{n:05d}", {"type": "line", "line": n}, "x" * 1000)
                 for n in range(1, 31)]
        self.assertEqual(agy_exam_index.index_settings({}), (12000, 3))
        for budget in (4000, 12000, 24000, 30000):
            resolved, workers = agy_exam_index.index_settings({"exam_index_batch_chars": budget})
            self.assertEqual(workers, 3)
            self.assertEqual(len(agy_exam_index._unit_batches(units, resolved)), ceil(30000 / budget))
        for key in ("exam_index_batch_chars", "exam_index_workers"):
            for invalid in (0, -1, True, 1.5, "3"):
                with self.assertRaisesRegex(ValueError, key):
                    agy_exam_index.index_settings({key: invalid})

    def test_engine_json_budget_controls_real_source_batches(self):
        config = Path(self.temporary.name) / "engine.json"
        source = self.questions / "large.txt"
        source.write_text("\n".join("x" * 1000 for _ in range(30)), encoding="utf-8")
        config.write_text(json.dumps({"exam_index_batch_chars": 10000, "exam_index_workers": 1}), encoding="utf-8")
        def covered(prompt, _schema, **_options):
            units = json.loads(prompt.split("Evidence units: ", 1)[1])
            return {"covered_unit_ids": [unit["id"] for unit in units if unit["coverage"] == "required"], "questions": []}
        with patch.dict(os.environ, {"TRANSCRIBER_CONFIG_PATH": str(config)}), patch(
            "agy_exam_index.agy_writer.request_json", side_effect=covered,
        ) as request:
            index = agy_exam_index.build_index(self.module, [source])
            self.assertEqual(request.call_count, 3)
            self.assertEqual(index["failed_sources"], [])
            config.write_text(json.dumps({"exam_index_workers": 0}), encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "exam_index_workers"):
                agy_exam_index.build_index(self.module, [source])

    def test_fake_agy_pool_speed_order_cache_and_incremental_publication(self):
        import time

        import cancellation
        from run_transcription import _publish_exam_index

        files = [self.questions / f"{n}.txt" for n in range(3)]
        for source in files:
            source.write_text("1. Explain dehydration.\n2. Explain fever.\n3. Explain shock.\n", encoding="utf-8")

        def fake(prompt, _schema, **_options):
            cancellation.run([sys.executable, "-c", "import time; time.sleep(.15)"], timeout=5)
            source = next(path for path in files if f"Source file: {path.name}\n" in prompt)
            units = json.loads(prompt.split("Evidence units: ", 1)[1])
            core = [unit for unit in units if unit["coverage"] == "required"]
            questions = [self._written_response(source, unit_ids=[unit["id"]], stem=unit["text"][3:])["questions"][0]
                         for unit in core]
            return {"covered_unit_ids": [unit["id"] for unit in core], "questions": questions}

        def measure(workers):
            (self.questions / agy_exam_index.STATE_NAME).unlink(missing_ok=True)
            publications = []
            def publish(index):
                publications.append(len(index["sources"]))
                _publish_exam_index(self.module, files, index)
            started = time.monotonic()
            with patch("agy_exam_index.agy_writer.request_json", side_effect=fake):
                result = agy_exam_index.build_index(self.module, files,
                    config={"exam_index_workers": workers, "exam_index_batch_chars": 24}, publish=publish)
            elapsed = time.monotonic() - started
            self.assertEqual(publications, list(range(1, 4)))
            return result, elapsed

        sequential, serial_seconds = measure(1)
        concurrent, parallel_seconds = measure(3)
        self.assertEqual(sequential, concurrent)
        self.assertLess(parallel_seconds, serial_seconds * .65)
        with patch("agy_exam_index.agy_writer.request_json", side_effect=AssertionError("cache miss")):
            resumed = agy_exam_index.build_index(self.module, files, config={"exam_index_workers": 3, "exam_index_batch_chars": 24})
        self.assertEqual(concurrent, resumed)
        print(f"Fake agy: serial={serial_seconds:.3f}s parallel={parallel_seconds:.3f}s speedup={serial_seconds / parallel_seconds:.2f}x")
        with patch("agy_exam_index.agy_writer.request_json", side_effect=fake) as request:
            agy_exam_index.build_index(self.module, files, config={"exam_index_batch_chars": 12000})
        self.assertEqual(request.call_count, 3)

    def test_concurrent_batch_failure_keeps_other_papers_and_valid_batches(self):
        files = [self.questions / name for name in ("A.txt", "B.txt")]
        for source in files:
            source.write_text("Source note one.\nSource note two.\nSource note three.\n", encoding="utf-8")
        def fake(prompt, _schema, **_options):
            units = json.loads(prompt.split("Evidence units: ", 1)[1])
            ids = [unit["id"] for unit in units if unit["coverage"] == "required"]
            if "Source file: A.txt" in prompt and ids == ["L00002"]:
                return {"covered_unit_ids": [], "questions": []}
            return {"covered_unit_ids": ids, "questions": []}
        config = {"exam_index_batch_chars": 20, "exam_index_workers": 3}
        with patch("agy_exam_index.agy_writer.request_json", side_effect=fake):
            index = agy_exam_index.build_index(self.module, files, config=config)
        self.assertEqual([source["file"] for source in index["sources"]], ["B.txt"])
        self.assertEqual([source["file"] for source in index["failed_sources"]], ["A.txt"])
        with patch("agy_exam_index.agy_writer.request_json", return_value={"covered_unit_ids": ["L00002"], "questions": []}) as request:
            resumed = agy_exam_index.build_index(self.module, files, config=config)
        self.assertEqual(request.call_count, 1)
        self.assertEqual(len(resumed["sources"]), 2)
        self.assertEqual(resumed["failed_sources"], [])

    def test_interrupted_build_keeps_completed_paper_and_resumes_cached_batches(self):
        import threading

        import cancellation
        from run_transcription import build_exam_index

        fast, slow = self.questions / "A.txt", self.questions / "B.txt"
        for source in (fast, slow):
            source.write_text("1. Explain dehydration.\n", encoding="utf-8")
        event = threading.Event()
        def fake(prompt, _schema, **_options):
            if "Source file: B.txt" in prompt:
                cancellation.wait(30)
            return self._written_response(fast)
        def progress(_done, _total, message):
            if "completed" in message:
                event.set()
        with patch("agy_exam_index.agy_writer.request_json", side_effect=fake), cancellation.request_scope(event):
            with self.assertRaises(cancellation.OperationCancelled):
                build_exam_index(self.module, {"exam_index_workers": 2}, progress)
        published = json.loads((self.questions / "exam-index.json").read_text())
        self.assertEqual([source["file"] for source in published["sources"]], ["A.txt"])
        with patch("agy_exam_index.agy_writer.request_json", return_value=self._written_response(slow)) as request:
            index, _target = build_exam_index(self.module, {"exam_index_workers": 2})
        self.assertEqual(request.call_count, 1)
        self.assertEqual([source["file"] for source in index["sources"]], ["A.txt", "B.txt"])

    def test_cancelled_pool_reaps_every_running_fake_agy(self):
        import threading
        import time

        import cancellation

        files = [self.questions / f"cancel-{n}.txt" for n in range(3)]
        for source in files:
            source.write_text("Source note.\n", encoding="utf-8")
        event = threading.Event()
        pids = [Path(self.temporary.name) / f"pid-{n}" for n in range(3)]
        lock = threading.Lock()
        called = []
        def fake(*_args, **_options):
            with lock:
                marker = pids[len(called)]
                called.append(marker)
            cancellation.run([sys.executable, "-c",
                "import os,time,pathlib; pathlib.Path(" + repr(str(marker)) + ").write_text(str(os.getpid())); time.sleep(30)"], timeout=35)
            raise AssertionError("fake agy was not stopped")
        def cancel():
            until = time.monotonic() + 5
            while not all(path.exists() for path in pids) and time.monotonic() < until:
                time.sleep(.01)
            event.set()
        timer = threading.Thread(target=cancel)
        timer.start()
        started = time.monotonic()
        try:
            with patch("agy_exam_index.agy_writer.request_json", side_effect=fake), cancellation.request_scope(event):
                with self.assertRaises(cancellation.OperationCancelled):
                    agy_exam_index.build_index(self.module, files, config={"exam_index_workers": 3})
        finally:
            event.set()
            timer.join()
        self.assertLess(time.monotonic() - started, 6)
        for path in pids:
            self.assertTrue(path.exists())
            with self.assertRaises(ProcessLookupError):
                os.kill(int(path.read_text()), 0)

    def test_index_preserves_source_answer_explanation_and_file_location(self):
        source = self.questions / "2023.txt"
        source.write_text(
            "1. Which option is printed as correct?\n"
            "A. Option A text\n"
            "B. Option B text\n"
            "Correct answer: B. Option B text\n"
            "Explanation: This explanation is printed in the paper.\n",
            encoding="utf-8",
        )
        units = agy_exam_index.read_source_units(source)
        by_text = {unit.text: unit.id for unit in units}
        expected = {
            "covered_unit_ids": [unit.id for unit in units],
            "questions": [{
                "unit_ids": [by_text["1. Which option is printed as correct?"],
                             by_text["A. Option A text"], by_text["B. Option B text"]],
                "number": 1,
                "kind": "mcq",
                "stem": "Which option is printed as correct?",
                "options": [{"label": "A", "text": "Option A text"}, {"label": "B", "text": "Option B text"}],
                "correct_option": "B",
                "answer_text": "Option B text",
                "answer_evidence": "Correct answer: B. Option B text",
                "answer_unit_ids": [by_text["Correct answer: B. Option B text"]],
                "explanation": "This explanation is printed in the paper.",
                "explanation_unit_ids": [by_text["Explanation: This explanation is printed in the paper."]],
                "section": None,
                "year": 2023,
                "year_evidence": "2023",
                "topic": None,
                "needs_review": False,
                "review_reason": None,
            }],
        }
        with patch("agy_exam_index.agy_writer.request_json", return_value=expected) as request:
            index = agy_exam_index.build_index(self.module, [source])
            write_index(index, self.questions)
            verified = paper_backed_index(index, {}, self.module.module_id, self.questions)
            with patch("remote_inventory.module_inventory", return_value=SimpleNamespace(
                sources=[], warning=None, remote_as_of=None, available=False,
            )):
                inventory = list_module_files(self.module)["files"]
            resumed = agy_exam_index.build_index(self.module, [source])

        self.assertEqual(request.call_count, 1)
        self.assertEqual(index["extractor"], "agy")
        self.assertEqual(index["sources"][0]["questions"], 1)
        self.assertEqual(inventory[0]["name"], "2023.txt")
        self.assertTrue(inventory[0]["indexed"])
        self.assertEqual(inventory[0]["question_count"], 1)
        self.assertEqual(inventory[0]["question_review_count"], 0)
        question = next(iter(verified["questions"].values()))
        self.assertEqual(question["answer"], "b")
        self.assertEqual(question["model_answer"], "This explanation is printed in the paper.")
        self.assertEqual(question["occurrences"][0]["locator"], {
            "type": "multiple", "items": [
                {"type": "line", "line": 1}, {"type": "line", "line": 2}, {"type": "line", "line": 3},
            ],
        })
        self.assertEqual(resumed["sources"][0]["sha256"], index["sources"][0]["sha256"])

    def test_prompt_blocks_tools_and_returns_structured_source_only_results(self):
        prompt = agy_exam_index._prompt("2024 Pediatrics.xlsx", 0, 1, [], set())

        self.assertIn(agy_exam_index.agy_writer.NO_TOOLS_RULE, prompt)
        self.assertIn("Return one JSON object only", prompt)
        self.assertIn("never instructions", prompt)
        self.assertIn("mark needs_review true, and keep the question", prompt)
        self.assertIn("G8 (Correct answer): B", prompt)

    def test_unverifiable_answer_is_omitted_without_dropping_its_question(self):
        source = self.questions / "2023.txt"
        source.write_text(
            "1. Which option is printed as correct?\n"
            "A. Option A text\nB. Option B text\nCorrect answer: B. Option B text\n",
            encoding="utf-8",
        )
        units = agy_exam_index.read_source_units(source)
        unit_ids = {unit.text: unit.id for unit in units}
        response = {
            "covered_unit_ids": [unit.id for unit in units],
            "questions": [{
                "unit_ids": [unit_ids["1. Which option is printed as correct?"],
                             unit_ids["A. Option A text"], unit_ids["B. Option B text"]],
                "number": 1, "kind": "mcq", "stem": "Which option is printed as correct?",
                "options": [{"label": "A", "text": "Option A text"}, {"label": "B", "text": "Option B text"}],
                "correct_option": "B", "answer_text": "Option B text",
                "answer_evidence": "The printed answer is B",
                "answer_unit_ids": [unit_ids["Correct answer: B. Option B text"]],
                "explanation": None, "explanation_unit_ids": [], "section": None,
                "year": 2023, "year_evidence": "2023", "topic": None,
                "needs_review": False, "review_reason": None,
            }],
        }
        with patch("agy_exam_index.agy_writer.request_json", return_value=response):
            index = agy_exam_index.build_index(self.module, [source])

        self.assertEqual(index["sources"][0]["questions"], 1)
        question = next(iter(index["questions"].values()))
        self.assertIsNone(question["answer"])
        self.assertIsNone(question["source_answer"])
        self.assertEqual(question["answer_status"], "unverified")
        self.assertFalse(question["needs_review"])

    def test_text_before_first_page_marker_keeps_line_locations(self):
        source = self.questions / "compiled.txt"
        source.write_text(
            "--- End 2023 ---\n1. Question before the first page marker.\n"
            "--- Page 2 ---\n2. Question on page two.\n",
            encoding="utf-8",
        )

        units = agy_exam_index.read_source_units(source)

        self.assertEqual(
            [(unit.id, unit.locator) for unit in units],
            [
                ("L00001", {"type": "line", "line": 1}),
                ("L00002", {"type": "line", "line": 2}),
                ("P00002", {"type": "page", "page": 2}),
            ],
        )
        self.assertEqual(units[1].text, "1. Question before the first page marker.")
        self.assertIn("2. Question on page two.", units[2].text)

    def test_spreadsheet_word_and_prepared_pdf_units_keep_document_locations(self):
        from docx import Document
        from openpyxl import Workbook

        spreadsheet = self.questions / "Exam.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Pediatrics"
        sheet.append(["Question", "A", "B"])
        sheet.append(["Which finding identifies the condition?", "Finding one", "Finding two"])
        workbook.save(spreadsheet)
        workbook.close()
        workbook_units = agy_exam_index.read_source_units(spreadsheet)
        self.assertEqual(workbook_units[1].locator, {
            "type": "spreadsheet_row", "sheet": "Pediatrics", "row": 2, "range": "A2:C2",
        })
        self.assertIn("A2 (Question): Which finding identifies the condition?", workbook_units[1].text)

        word = self.questions / "Exam.docx"
        document = Document()
        document.add_paragraph("Question text in a paragraph.")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "A. First option"
        table.cell(0, 1).text = "B. Second option"
        document.save(word)
        word_units = agy_exam_index.read_source_units(word)
        self.assertEqual(word_units[0].locator, {"type": "paragraph", "paragraph": 1})
        self.assertEqual(word_units[1].locator, {"type": "table_row", "table": 2, "row": 1})

        pdf = self.questions / "Exam.pdf"
        pdf.write_bytes(b"prepared in a separate OCR step")
        prepared = self.questions / "Exam.pdf.txt"
        prepared.write_text("--- Page 2 ---\nQuestion on the second page.\n", encoding="utf-8")
        pdf_units = agy_exam_index.read_source_units(pdf, prepared)
        self.assertEqual(pdf_units[0].locator, {"type": "page", "page": 2})
        self.assertEqual(pdf_units[0].text, "Question on the second page.")

    def test_spreadsheet_row_years_verify_past_exam_badges(self):
        from exam_preparation import prepare_exam_file
        from openpyxl import Workbook
        from question_provenance import assessment_catalog, assessment_verified_years
        from run_transcription import assessment_entry

        bank = self.questions / "Infection questions.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Bank"
        sheet.append(["Text", "A", "B", "Correct", "Year"])
        sheet.append(["Which organism causes the infection?", "Organism one", "Organism two", "B", 2021])
        sheet.append(["Which drug treats the infection?", "Drug one", "Drug two", "A", None])
        workbook.save(bank)
        workbook.close()
        self.assertEqual(prepare_exam_file(self.module, "Questions/Infection questions.xlsx")["status"], "ready")
        units = agy_exam_index.read_source_units(bank)

        def row_question(unit, stem, options, answer, year):
            return {
                "unit_ids": [unit.id], "number": None, "kind": "mcq", "stem": stem,
                "options": [{"label": label, "text": text} for label, text in zip("AB", options)],
                "correct_option": answer, "answer_text": dict(zip("AB", options))[answer],
                "answer_evidence": f"D{unit.locator['row']} (Correct): {answer}", "answer_unit_ids": [unit.id],
                "explanation": None, "explanation_unit_ids": [], "section": None, "year": year,
                "year_evidence": None if year is None else f"E{unit.locator['row']} (Year): {year}",
                "topic": None, "needs_review": False, "review_reason": None,
            }

        response = {"covered_unit_ids": [unit.id for unit in units], "questions": [
            row_question(units[1], "Which organism causes the infection?", ("Organism one", "Organism two"), "B", 2021),
            row_question(units[2], "Which drug treats the infection?", ("Drug one", "Drug two"), "A", None),
        ]}
        with patch("agy_exam_index.agy_writer.request_json", return_value=response):
            write_index(agy_exam_index.build_index(self.module, [bank]), self.questions)
        occurrences = [occurrence for question in json.loads((self.questions / "exam-index.json").read_text(encoding="utf-8"))["questions"].values()
                       for occurrence in question["occurrences"]]
        self.assertEqual(sorted((occurrence["year"] or 0, occurrence["year_source"] or "") for occurrence in occurrences),
                         [(0, ""), (2021, "question")])
        catalog = assessment_catalog(self.root, {"assessment_sources": [assessment_entry(bank)]})
        self.assertEqual(assessment_verified_years(catalog), {2021})
        self.assertEqual(assessment_verified_years(assessment_catalog(self.root, {"assessment_sources": []})), set())

    def test_adding_a_paper_rebuilds_only_that_paper(self):
        from agy_index_fixtures import _proposal

        def paper(name, stem):
            path = self.questions / name
            path.write_text(f"Final examination {name[-8:-4]}\n1. {stem}?\na. Fever\nb. Rash\nAnswer: a\n", encoding="utf-8")
            return path

        papers = [paper("Final 2021.txt", "Which sign comes first in measles"), paper("Final 2022.txt", "Which sign marks scarlet fever")]
        calls, published, events = [], [], []

        def answer(prompt, schema, **options):
            calls.append("review" if "Review questions: " in prompt else prompt.split("Source file: ", 1)[1].splitlines()[0])
            return _proposal({path.name: path for path in self.questions.iterdir()})(prompt, schema, **options)

        def run(files):
            calls.clear(), published.clear(), events.clear()
            # The class stubs the review; this test needs its saved cache record.
            with patch("agy_exam_index.agy_writer.request_json", side_effect=answer), patch(
                "exam_question_review.review_paper", REVIEW_PAPER,
            ):
                return agy_exam_index.build_index(
                    self.module, files, config={}, publish=published.append,
                    report_progress=lambda done, total, message, **details: events.append((details.get("file"), details.get("phase"))))

        first = run(papers)
        self.assertEqual(len(published), 2)
        state_path = self.questions / agy_exam_index.STATE_NAME
        state_before = state_path.stat().st_mtime_ns
        papers.append(paper("Final 2023.txt", "Which rash spares the face"))
        second = run(papers)
        self.assertEqual([call for call in calls if call != "review"], ["Final 2023.txt"])
        self.assertEqual(len(published), 1)
        self.assertEqual({name for name, _phase in events}, {"Final 2021.txt", "Final 2022.txt", "Final 2023.txt"})
        self.assertEqual([phase for name, phase in events if name != "Final 2023.txt"], ["done", "done"])
        self.assertIn("extracting", [phase for name, phase in events if name == "Final 2023.txt"])
        self.assertEqual([source["file"] for source in second["sources"]], ["Final 2021.txt", "Final 2022.txt", "Final 2023.txt"])
        self.assertEqual(second["sources"][:2], first["sources"])
        self.assertGreater(state_path.stat().st_mtime_ns, state_before)

        papers[0].write_text(papers[0].read_text(encoding="utf-8").replace("measles", "rubella"), encoding="utf-8")
        run(papers)
        self.assertEqual([call for call in calls if call != "review"], ["Final 2021.txt"])
        self.assertIn("extracting", [phase for name, phase in events if name == "Final 2021.txt"])
        self.assertEqual([phase for name, phase in events if name == "Final 2022.txt"], ["done"])

    def test_each_required_batch_is_saved_and_reused_on_retry(self):
        source = self.questions / "questions.txt"
        source.write_text("\n".join(f"Source note {number}: " + "evidence " * 25 for number in range(3)), encoding="utf-8")

        def covered(prompt, _schema, **_options):
            evidence = json.loads(prompt.split("Evidence units: ", 1)[1])
            return {
                "covered_unit_ids": [unit["id"] for unit in evidence if unit["coverage"] == "required"],
                "questions": [],
            }

        with patch("agy_exam_index.index_settings", return_value=(260, 1)), patch(
            "agy_exam_index.agy_writer.request_json", side_effect=covered,
        ) as request:
            agy_exam_index.build_index(self.module, [source])
            agy_exam_index.build_index(self.module, [source])

        self.assertEqual(request.call_count, 3)

    def test_large_spreadsheet_batches_bound_questions_and_prompt_size(self):
        units = [
            agy_exam_index.SourceUnit(
                f"S001R{row:06d}",
                {"type": "spreadsheet_row", "sheet": "Pediatrics", "row": row,
                 "range": f"A{row}:AF{row}"},
                f"Question {row}: " + "source evidence " * 48,
            )
            for row in range(1, 504)
        ]

        batches = agy_exam_index._unit_batches(units, 4_000)
        prompt_sizes = [
            len(agy_exam_index._prompt(
                "Pediatrics.xlsx", index, len(batches), batch,
                {unit.id for unit in core},
            ).encode("utf-8"))
            for index, (core, batch) in enumerate(batches)
        ]

        self.assertGreater(len(batches), 1)
        self.assertEqual([unit.id for core, _batch in batches for unit in core],
                         [unit.id for unit in units])
        self.assertLessEqual(max(len(core) for core, _batch in batches), 5)
        self.assertLessEqual(max(prompt_sizes), 20_000)

    def test_unquoted_model_explanation_is_omitted_without_dropping_question(self):
        source = self.questions / "2023.txt"
        source.write_text(
            "1. Which option is printed as correct?\n"
            "A. Option A text\nB. Option B text\nCorrect answer: B. Option B text\n",
            encoding="utf-8",
        )
        units = agy_exam_index.read_source_units(source)
        unit_ids = {unit.text: unit.id for unit in units}
        response = {
            "covered_unit_ids": [unit.id for unit in units],
            "questions": [{
                "unit_ids": [unit_ids["1. Which option is printed as correct?"],
                             unit_ids["A. Option A text"], unit_ids["B. Option B text"]],
                "number": 1, "kind": "mcq", "stem": "Which option is printed as correct?",
                "options": [{"label": "A", "text": "Option A text"}, {"label": "B", "text": "Option B text"}],
                "correct_option": "B", "answer_text": "Option B text",
                "answer_evidence": "Correct answer: B. Option B text",
                "answer_unit_ids": [unit_ids["Correct answer: B. Option B text"]],
                "explanation": "The incorrect option causes an unrelated medical condition.",
                "explanation_unit_ids": [unit_ids["1. Which option is printed as correct?"]],
                "section": None, "year": 2023, "year_evidence": "2023", "topic": None,
                "needs_review": False, "review_reason": None,
            }],
        }
        with patch("agy_exam_index.agy_writer.request_json", return_value=response):
            index = agy_exam_index.build_index(self.module, [source])

        question = next(iter(index["questions"].values()))
        self.assertEqual(index["sources"][0]["questions"], 1)
        self.assertEqual(question["model_answer"], "")
        self.assertIsNone(question["occurrences"][0]["explanation"])

    def test_literal_source_id_expression_and_string_year_are_normalized(self):
        source = self.questions / "2023.txt"
        source.write_text(
            "1. Which option is printed as correct?\n"
            "A. Option A text\nB. Option B text\nCorrect answer: B. Option B text\n",
            encoding="utf-8",
        )
        units = agy_exam_index.read_source_units(source)
        unit_ids = {unit.text: unit.id for unit in units}
        response = {
            "covered_unit_ids": [unit.id for unit in units],
            "questions": [{
                "unit_ids": [unit_ids["1. Which option is printed as correct?"],
                             unit_ids["A. Option A text"], unit_ids["B. Option B text"]],
                "number": 1, "kind": "mcq", "stem": "Which option is printed as correct?",
                "options": [{"label": "A", "text": "Option A text"}, {"label": "B", "text": "Option B text"}],
                "correct_option": "B", "answer_text": "Option B text",
                "answer_evidence": "Correct answer: B. Option B text",
                "answer_unit_ids": [unit_ids["Correct answer: B. Option B text"]],
                "explanation": None, "explanation_unit_ids": [], "section": None,
                "year": "2023", "year_evidence": "2023", "topic": None,
                "needs_review": False, "review_reason": None,
            }],
        }
        malformed = json.dumps(response).replace(
            '"unit_ids": ["L00001", "L00002"',
            '"unit_ids": ["L00001", "L00001".replace("1", "2")',
            1,
        )
        proposal_error = agy_exam_index.agy_writer.AgyProposalError("invalid JSON", malformed)

        with patch("agy_exam_index.agy_writer.request_json", side_effect=proposal_error):
            index = agy_exam_index.build_index(self.module, [source])

        question = next(iter(index["questions"].values()))
        self.assertEqual(index["sources"][0]["questions"], 1)
        self.assertEqual(question["years"], [2023])
        self.assertEqual(question["stem"], "Which option is printed as correct?")
        self.assertEqual(question["answer"], "b")

    def test_invalid_json_is_retried_against_the_same_source_units(self):
        source = self.questions / "questions.txt"
        source.write_text("1. Explain dehydration.\n", encoding="utf-8")
        unit = agy_exam_index.read_source_units(source)[0]
        response = {
            "covered_unit_ids": [unit.id],
            "questions": [{
                "unit_ids": [unit.id], "number": 1, "kind": "written", "stem": "Explain dehydration.",
                "options": [], "correct_option": None, "answer_text": None, "answer_evidence": None,
                "answer_unit_ids": [], "explanation": None, "explanation_unit_ids": [], "section": None,
                "year": None, "year_evidence": None, "topic": None,
                "needs_review": True, "review_reason": "The source does not print an answer.",
            }],
        }
        malformed = agy_exam_index.agy_writer.AgyProposalError("invalid JSON", '{"questions": [{"broken" "field"}]}')

        with patch("agy_exam_index.agy_writer.request_json", side_effect=[malformed, response]) as writer:
            index = agy_exam_index.build_index(self.module, [source], config={"exam_index_request_timeout_seconds": 137})
        self.assertEqual([call.kwargs["timeout"] for call in writer.call_args_list], [137, 137])

        question = next(iter(index["questions"].values()))
        self.assertEqual(index["sources"][0]["questions"], 1)
        self.assertEqual(question["stem"], "Explain dehydration.")
        self.assertFalse(question["needs_review"])
        self.assertEqual(question["answer_status"], "no_key")

    def test_unknown_model_falls_back_and_records_producing_model(self):
        source = self.questions / "paper.txt"
        source.write_text("Explain dehydration.")
        with patch("agy_exam_index.agy_writer.request_json", side_effect=[
            agy_exam_index.agy_writer.AgyWriterError('invalid model selection: model low is not recognized'),
            self._written_response(source),
        ]) as writer:
            index = agy_exam_index.build_index(self.module, [source], config={"exam_index_models": ["low", "high"]})
        self.assertEqual(index["failed_sources"], [])
        self.assertEqual([call.kwargs["model"] for call in writer.call_args_list], ["low", "high"])
        payload = json.loads((self.questions / agy_exam_index.STATE_NAME).read_text())["files"][source.name]["batches"]["0"]["payload"]
        self.assertEqual((payload["model"], payload["attempts"]), ("high", 2))
        self.backoff.assert_not_called()

    def test_transient_retry_preserves_batch_and_uses_backoff(self):
        source = self.questions / "paper.txt"
        source.write_text("Explain dehydration.")
        with patch("agy_exam_index.agy_writer.request_json", side_effect=[
            agy_exam_index.agy_writer.AgyWriterError("HTTP 503: network request failed"),
            self._written_response(source),
        ]) as writer:
            index = agy_exam_index.build_index(self.module, [source], config={})
        self.assertEqual(index["failed_sources"], [])
        self.assertEqual(writer.call_count, 2)
        self.assertEqual(writer.call_args_list[0], writer.call_args_list[1])
        self.backoff.assert_called_once_with(15)

    def test_repair_phase_runs_only_failed_batch_and_reuses_saved_payloads(self):
        source = self.questions / "paper.txt"
        source.write_text("Explain dehydration.\nExplain fever.")
        calls = []
        failed = False

        def writer(prompt, _schema, **_options):
            nonlocal failed
            units = json.loads(prompt.split("Evidence units: ", 1)[1])
            core = [unit for unit in units if unit["coverage"] == "required"]
            calls.append(core[0]["id"])
            if core[0]["id"] == "L00002" and not failed:
                failed = True
                raise agy_exam_index.agy_writer.AgyWriterError("HTTP 502")
            return {"covered_unit_ids": [unit["id"] for unit in core], "questions": [
                self._written_response(source, stem=core[0]["text"], unit_ids=[core[0]["id"]])["questions"][0]]}

        events = []
        with patch("agy_exam_index.agy_writer.request_json", side_effect=writer):
            index = agy_exam_index.build_index(self.module, [source], config={
                "exam_index_workers": 1, "exam_index_batch_chars": 20, "exam_index_retry_limit": 0,
            }, report_progress=lambda _done, _total, _message, **details: events.append(details))
        self.assertEqual(index["failed_sources"], [])
        self.assertEqual(calls, ["L00001", "L00002", "L00002"])
        self.assertEqual([event["file"] for event in events if event["phase"] == "repairing"], [source.name])
        self.assertFalse(any(event["phase"] == "failed" for event in events))

    def test_pdf_text_repair_preserves_original_and_runs_once(self):
        from atomic_io import _atomic_write_text

        source = self.questions / "paper.pdf"
        source.write_bytes(b"original PDF")
        prepared = self.questions / "paper.pdf.txt"
        prepared.write_text("Unreadable prepared text")
        calls = []

        def repair(_module, paper):
            calls.append(paper.name)
            _atomic_write_text(agy_exam_index._ocr_text_path(self.module, paper), "Explain dehydration.")

        def writer(prompt, _schema, **_options):
            units = json.loads(prompt.split("Evidence units: ", 1)[1])
            return {"covered_unit_ids": [unit["id"] for unit in units], "questions": []}

        with patch("exam_preparation.exam_file_status", return_value={"preparation": "ready"}), patch(
            "agy_exam_index._repair_pdf_text", side_effect=repair
        ), patch("agy_exam_index.agy_writer.request_json", side_effect=writer) as request:
            index = agy_exam_index.build_index(self.module, [source], config={})
        self.assertEqual(calls, [source.name])
        self.assertEqual(request.call_count, 2)
        self.assertEqual(index["failed_sources"][0]["cause"], "source_unreadable")
        self.assertEqual(prepared.read_text(), "Unreadable prepared text")
        self.assertEqual(source.read_bytes(), b"original PDF")

    def test_fragmented_pdf_text_layer_is_ocr_repaired_before_indexing(self):
        from atomic_io import _atomic_write_text

        source = self.questions / "paper.pdf"
        source.write_bytes(b"original PDF")
        (self.questions / "paper.pdf.txt").write_text("1)Cl\n   assi\n      fi\n" * 20)
        prompts = []

        def repair(_module, paper):
            _atomic_write_text(agy_exam_index._ocr_text_path(self.module, paper), "1) Classification of poisoning.")

        def writer(prompt, _schema, **_options):
            prompts.append(prompt)
            units = json.loads(prompt.split("Evidence units: ", 1)[1])
            return {"covered_unit_ids": [unit["id"] for unit in units], "questions": []}

        with patch("exam_preparation.exam_file_status", return_value={"preparation": "ready"}), patch(
            "agy_exam_index._repair_pdf_text", side_effect=repair
        ) as repaired, patch("agy_exam_index.agy_writer.request_json", side_effect=writer):
            agy_exam_index.build_index(self.module, [source], config={})
        repaired.assert_called_once()
        self.assertTrue(prompts)
        self.assertFalse(any("1)Cl" in prompt for prompt in prompts))
        self.assertTrue(any("Classification of poisoning" in prompt for prompt in prompts))

    def test_pdf_stronger_bilingual_ocr_recovers_zero_or_unsupported_extraction(self):
        for unsupported in (False, True):
            with self.subTest(unsupported=unsupported):
                source = self.questions / ("unsupported.pdf" if unsupported else "empty.pdf")
                source.write_bytes(source.name.encode())
                prepared = self.questions / (source.name + ".txt")
                prepared.write_text("Explain dehydration.")
                initial = self._written_response(prepared) if unsupported else {"covered_unit_ids": ["L00001"], "questions": []}
                if unsupported:
                    initial["questions"].extend([
                        self._written_response(prepared, stem="Invented stem from nowhere")["questions"][0]
                        for _ in range(2)
                    ])
                repaired = self._written_response(prepared, stem="Explain renal injury.")
                with patch("exam_preparation.exam_file_status", return_value={"preparation": "ready"}), patch(
                    "source_preparation._ocr_pdf"
                ) as ocr, patch("question_coverage.pdf_text", return_value="Explain renal injury."), patch(
                    "agy_exam_index.agy_writer.request_json", side_effect=[initial, repaired]
                ) as request:
                    index = agy_exam_index.build_index(self.module, [source], config={})
                    with patch("agy_exam_index.agy_writer.request_json", side_effect=AssertionError("OCR batches should be cached")):
                        cached = agy_exam_index.build_index(self.module, [source], config={})
                self.assertEqual(index, cached)
                self.assertEqual(index["failed_sources"], [])
                self.assertEqual(request.call_count, 2)
                self.assertEqual([call.kwargs["page_segmentation"] for call in ocr.call_args_list], [None, 6])
                self.assertEqual(ocr.call_args.args[0], source)
                self.assertEqual(ocr.call_args.args[2], "ara+eng")
                self.assertTrue(ocr.call_args.kwargs["stronger"])
                self.assertTrue(index["sources"][0]["prepared_file"].startswith(".exam-index-ocr-"))
                self.assertEqual(prepared.read_text(), "Explain dehydration.")
                self.assertEqual(len(paper_backed_index(index, {}, self.module.module_id, self.questions)["questions"]), 1)

    def test_pdf_text_repair_keeps_the_ocr_reading_with_more_question_stems(self):
        source = self.questions / "paper.pdf"
        source.write_bytes(b"highlighted scan")
        readings = ["1) Kept stem.\n2) Another.\n", "1) Kept stem.\n2) Another.\n3) Highlighted stem.\n"]
        with patch("source_preparation._ocr_pdf") as ocr, patch("question_coverage.pdf_text", side_effect=readings):
            agy_exam_index._repair_pdf_text(self.module, source)
        self.assertEqual(ocr.call_count, 2)
        self.assertIn("Highlighted stem", agy_exam_index._ocr_text_path(self.module, source).read_text())

    def test_build_wrapper_reuses_ocr_after_normal_preparation_failure(self):
        from run_transcription import build_exam_index

        source = self.questions / "paper.pdf"
        source.write_bytes(b"original scan")
        prepared = self.questions / "paper.pdf.txt"
        prepared.write_text("Unusable original text")
        preparation_state = self.root / ".transcriber-cache" / "exam-texts.json"
        preparation_state.parent.mkdir(parents=True)
        preparation_state.write_text(json.dumps({source.name: {
            "sha256": agy_exam_index._sha256(source), "text": prepared.name,
            "prepared": True, "error": "No readable text",
        }}))
        original_units = agy_exam_index._unit_text("Explain dehydration.")
        response = {"covered_unit_ids": [original_units[0].id],
                    "questions": [self._written_response(prepared, stem="Explain dehydration.")["questions"][0]]}
        with patch("exam_preparation.prepare_exam_file", return_value={"status": "failed", "message": "No readable text"}) as prepare, patch(
            "exam_preparation.exam_file_status", return_value={"preparation": "failed"}
        ), patch("source_preparation._ocr_pdf") as ocr, patch("question_coverage.pdf_text", return_value="Explain dehydration."), patch(
            "agy_exam_index.agy_writer.request_json", return_value=response
        ) as writer:
            first, target = build_exam_index(self.module, {})
            second, _target = build_exam_index(self.module, {})
        self.assertEqual(first, second)
        self.assertEqual(first["failed_sources"], [])
        self.assertTrue(target.is_file())
        self.assertEqual((prepare.call_count, ocr.call_count, writer.call_count), (1, len(agy_exam_index.OCR_PAGE_SEGMENTATIONS), 1))
        self.assertEqual(prepared.read_text(), "Unusable original text")

    def test_changed_pdf_during_ocr_does_not_publish_repaired_receipt(self):
        source = self.questions / "paper.pdf"
        source.write_bytes(b"original scan")
        destination = agy_exam_index._ocr_text_path(self.module, source)

        def replace(_source, _destination, _language, **_options):
            source.write_bytes(b"changed original")

        with patch("source_preparation._ocr_pdf", side_effect=replace), patch("question_coverage.pdf_text", return_value="Explain dehydration."):
            with self.assertRaisesRegex(agy_exam_index.SourceChangedError, "changed during OCR"):
                agy_exam_index._repair_pdf_text(self.module, source)
        self.assertFalse(destination.exists())
        self.assertFalse(agy_exam_index._ocr_text_path(self.module, source).exists())

    def test_pdf_zero_questions_without_tools_has_one_extraction_attempt(self):
        source = self.questions / "paper.pdf"
        source.write_bytes(b"original PDF")
        (self.questions / "paper.pdf.txt").write_text("Readable but unrecognized source text")
        with patch("exam_preparation.exam_file_status", return_value={"preparation": "ready"}), patch(
            "source_preparation.shutil.which", return_value=None
        ), patch("agy_exam_index.agy_writer.request_json", return_value={"covered_unit_ids": ["L00001"], "questions": []}) as request:
            index = agy_exam_index.build_index(self.module, [source], config={})
        self.assertEqual(request.call_count, 1)
        self.assertEqual(index["failed_sources"][0]["cause"], "source_ocr_unavailable")
        self.assertEqual(index["failed_sources"][0]["attempts"], 1)

    def test_pdf_without_ocr_tools_reports_plain_source_reason(self):
        source = self.questions / "paper.pdf"
        source.write_bytes(b"original PDF")
        (self.questions / "paper.pdf.txt").write_text("??? 1234 !!!")
        with patch("exam_preparation.exam_file_status", return_value={"preparation": "ready"}), patch(
            "source_preparation.shutil.which", return_value=None
        ), patch("agy_exam_index.agy_writer.request_json") as request:
            index = agy_exam_index.build_index(self.module, [source], config={})
        request.assert_not_called()
        self.assertEqual(index["failed_sources"][0]["cause"], "source_ocr_unavailable")
        self.assertEqual(index["failed_sources"][0]["attempts"], 0)
        self.assertIn("Arabic and English OCR tools are unavailable", index["failed_sources"][0]["error"])

    def test_cancellation_interrupts_retry_backoff(self):
        from threading import Event

        import cancellation
        from exam_index_recovery import request_with_retry, retry_settings

        cancelled = Event()

        def backoff(seconds):
            cancelled.set()
            cancellation.wait(seconds)

        self.backoff.side_effect = backoff
        with cancellation.request_scope(cancelled), patch("agy_exam_index.agy_writer.request_json", side_effect=agy_exam_index.agy_writer.AgyWriterError("503")) as request:
            with self.assertRaises(cancellation.OperationCancelled):
                request_with_retry(lambda model: agy_exam_index.agy_writer.request_json("prompt", {}, model=model), retry_settings({}))
        self.assertEqual(request.call_count, 1)

    def test_filename_year_examples_flow_through_occurrences_and_duplicate_collapse(self):
        import copy

        from exam_index import _collapse_duplicates
        from exam_years import extract_filename_exam_years

        examples = {"End 2021.pdf": 2021, "End Toxico 2023.pdf": 2023,
                    "Final 21.pdf": 2021, "Final 25.pdf": 2025,
                    "فاينال توكسو 2024.pdf": 2024, "فاينال توكسو ٢٠٢٤.pdf": 2024,
                    "CamScanner ٢٩-١١-٢٠٢٣ ١٤.٤٧.pdf": None,
                    "MCQs integrated + 2018_Trim.pdf": 2018,
                    "Khalsa questions of toxo.pdf": None, "acetyl salicylic acid MCQ.pdf": None}
        for filename, year in examples.items():
            with self.subTest(filename=filename):
                self.assertEqual(extract_filename_exam_years(filename), (year,) if year else ())
                source = self.questions / filename
                source.write_bytes(b"PDF original")
                prepared = self.questions / (filename + ".txt")
                prepared.write_text("Explain dehydration.")
                with patch("exam_preparation.exam_file_status", return_value={"preparation": "ready"}):
                    snapshot = agy_exam_index._source_snapshot(self.module, source, 12000)
                response = {"covered_unit_ids": ["L00001"], "questions": [self._written_response(prepared)["questions"][0]]}
                questions, dropped = agy_exam_index._validate_batch(response, snapshot.units, {"L00001"}, snapshot)
                record = agy_exam_index._source_record(snapshot, questions, dropped)
                from exam_index import index_from_sources
                index = index_from_sources(self.module.module_id, [record])
                question = next(iter(index["questions"].values()))
                self.assertEqual(question["years"], [year] if year else [])
                self.assertEqual(question["occurrences"][0]["year_source"], "file_name" if year else None)
                self.assertEqual(question["answer_status"], "no_key")
                self.assertIsNone(question["answer"])
                index["questions"]["duplicate"] = copy.deepcopy(question)
                _collapse_duplicates(index["questions"])
                self.assertEqual(next(iter(index["questions"].values()))["years"], [year] if year else [])
                verified = paper_backed_index(index, {}, self.module.module_id, self.questions)
                self.assertEqual(len(verified["questions"]), 1)

    def test_empty_cli_output_is_transient_and_retries_without_splitting(self):
        source = self.questions / "paper.txt"
        source.write_text("Explain dehydration.")
        error = agy_exam_index.agy_writer.AgyWriterError("agy returned invalid JSON: Expecting value: line 1 column 1 (char 0)")
        self.assertFalse(agy_exam_index._capacity_error(error))
        with patch("agy_exam_index.agy_writer.request_json", side_effect=[error, self._written_response(source)]) as writer:
            index = agy_exam_index.build_index(self.module, [source], config={})
        self.assertEqual(index["failed_sources"], [])
        self.assertEqual(writer.call_count, 2)
        self.assertEqual(writer.call_args_list[0], writer.call_args_list[1])
        self.backoff.assert_called_once_with(15)

    def test_recovery_config_rejects_invalid_json_settings(self):
        from exam_index_recovery import retry_settings

        invalid = {"exam_index_retry_limit": [-1, True, "2", 1.5],
                   "exam_index_retry_backoff_seconds": [[], [True], [-1], [float("nan")], "15"],
                   "exam_index_models": [[], [""], [3], "low"]}
        for key, values in invalid.items():
            for value in values:
                with self.subTest(key=key, value=value), self.assertRaisesRegex(ValueError, key):
                    retry_settings({key: value})

    def _written_response(self, source, **overrides):
        units = agy_exam_index.read_source_units(source)
        question = {
            "unit_ids": [unit.id for unit in units], "number": 1, "kind": "written",
            "stem": "Explain dehydration.", "options": [], "correct_option": None,
            "answer_text": None, "answer_evidence": None, "answer_unit_ids": [],
            "explanation": None, "explanation_unit_ids": [], "section": None,
            "year": None, "year_evidence": None, "topic": None,
            "needs_review": False, "review_reason": None,
        }
        question.update(overrides)
        return {"covered_unit_ids": [unit.id for unit in units], "questions": [question]}

    def test_advisory_review_distinguishes_missing_keys_from_defects_on_cache_replay(self):
        source = self.questions / "paper.txt"
        source.write_text("1. Explain dehydration.\n", encoding="utf-8")
        for reason, review in [
            ("No printed answer key provided in the source text.", False),
            ("No explicit answer key is available in the source.", False),
            ("No explicit answer key in the source text.", False),
            ("No explicit answer key or marked correct option is provided in the source text.", False),
            ("OCR damage makes the text unreadable.", True),
            (".", True),
            ("No printed answer key; handwritten marks are ambiguous.", True),
            ("No explicit answer key or marked correct option found in the provided units.", False),
            ("No answer key or marked correct option provided for short notes question in source text.", False),
            ("Fill-in-the-blank question without an official printed answer key.", False),
            ("No official answer key.", False),
            ("No explicit answer key or marked correct option in the source text; option 'a' missing.", True),
            ("No explicit answer key or marked correct option in printed text; option (a) label is missing in OCR.", True),
            ("Stem is incomplete/truncated in the source text; no answer key provided.", True),
            ("True/false question; contains handwritten marks in parentheses but no official printed key.", True),
            ("No official answer key; student annotations present.", True),
            ("Written completion question; no official answer key provided in source text.", False),
            ("Case clinical sub-question with student handwriting; no official answer key provided in source text.", True),
            ("The stem is a fragment of another question.", True),
        ]:
            with self.subTest(reason=reason):
                response = self._written_response(source, needs_review=True, review_reason=reason)
                state_path = self.questions / agy_exam_index.STATE_NAME
                state_path.unlink(missing_ok=True)
                with patch("agy_exam_index.agy_writer.request_json", return_value=response):
                    index = agy_exam_index.build_index(self.module, [source])
                with patch("agy_exam_index.agy_writer.request_json", side_effect=AssertionError("cache miss")):
                    replay = agy_exam_index.build_index(self.module, [source])
                self.assertEqual(replay, index)
                question = next(iter(index["questions"].values()))
                self.assertEqual(question["needs_review"], review)
                self.assertEqual(question["answer_status"], "no_key")
                self.assertIsNone(question["answer"])
                self.assertEqual(question["occurrences"][0]["model_review_reason"], reason)
                self.assertEqual(index["sources"][0]["needs_review"], int(review))

    def test_ungrounded_topic_is_dropped_without_flagging_the_question(self):
        source = self.questions / "paper.txt"
        source.write_text("1. Explain dehydration.\n", encoding="utf-8")
        response = self._written_response(source, topic="Cardiology", needs_review=True,
                                          review_reason="No explicit answer key in the source text.")
        with patch("agy_exam_index.agy_writer.request_json", return_value=response):
            index = agy_exam_index.build_index(self.module, [source])
        question = next(iter(index["questions"].values()))
        self.assertIsNone(question["topic"])
        self.assertFalse(question["needs_review"])
        self.assertIn("topic is not grounded", question["occurrences"][0]["review_reason"])

    def test_unverifiable_answer_is_omitted_without_flagging_the_question(self):
        source = self.questions / "paper.txt"
        source.write_text("1. Choose the cause.\nA. Fluid loss\nB. Reduced intake\n", encoding="utf-8")
        response = self._written_response(
            source, kind="mcq", stem="Choose the cause.", correct_option="b", answer_unit_ids=["missing"],
            options=[{"label": "a", "text": "Fluid loss"}, {"label": "b", "text": "Reduced intake"}],
        )
        with patch("agy_exam_index.agy_writer.request_json", return_value=response):
            index = agy_exam_index.build_index(self.module, [source])
        question = next(iter(index["questions"].values()))
        self.assertIsNone(question["answer"])
        self.assertEqual(question["answer_status"], "unverified")
        self.assertFalse(question["needs_review"])
        self.assertEqual(index["sources"][0]["needs_review"], 0)

    def test_khalsa_chopped_words_retain_stem_options_and_printed_answer(self):
        source = self.questions / "Khalsa questions of toxo.txt"
        source.write_text(
            "Cl\nassi\nfi\ncat\nionofpoi\nsoni\nngaccor\ndingto\ntheir\nmodeof\nactioninto\n"
            "A. Lo\ncal\nB. Sys\ntemic\nCorrect ans\nwer: B Sys\ntemic\n", encoding="utf-8",
        )
        response = self._written_response(
            source, kind="mcq", stem="Classification of poisoning according to their mode of action into",
            options=[{"label": "a", "text": "Local"}, {"label": "b", "text": "Systemic"}],
            correct_option="b", answer_text="Systemic", answer_evidence="Correct answer: B Systemic",
            answer_unit_ids=[unit.id for unit in agy_exam_index.read_source_units(source)],
        )
        with patch("agy_exam_index.agy_writer.request_json", return_value=response):
            index = agy_exam_index.build_index(self.module, [source])
        question = next(iter(index["questions"].values()))
        self.assertEqual(question["answer_status"], "printed")
        self.assertEqual(question["options"], {"a": "Local", "b": "Systemic"})
        self.assertFalse(question["needs_review"])
        response["questions"][0]["stem"] = "Classificatlon of poisoning according to their mode of action into"
        (self.questions / agy_exam_index.STATE_NAME).unlink()
        with patch("agy_exam_index.agy_writer.request_json", return_value=response):
            approximate = agy_exam_index.build_index(self.module, [source])
        self.assertEqual(len(approximate["questions"]), 1)
        response["questions"][0]["stem"] = "Classification of poisoning according to imaginary invented mechanisms"
        (self.questions / agy_exam_index.STATE_NAME).unlink()
        with patch("agy_exam_index.agy_writer.request_json", return_value=response):
            rejected = agy_exam_index.build_index(self.module, [source])
        self.assertEqual(rejected["questions"], {})
        self.assertEqual(rejected["sources"][0]["dropped_question_reasons"], {"A question stem is not grounded": 1})

    def test_arabic_year_evidence_is_verified_and_mismatched_year_needs_review(self):
        cases = [
            ("paper.txt", "امتحان ٢٠٢٣", 2023),
            ("CamScanner ٢٩-١١-٢٠٢٣ ١٤.٤٧.txt", "CamScanner ٢٩-١١-٢٠٢٣ ١٤.٤٧.txt", 2023),
            ("paper.txt", "٢٠٢٣", 2023),
            ("paper.txt", "2019", 2023),
            ("paper.txt", "unprinted heading 2023", 2023),
        ]
        for filename, evidence, expected in cases:
            with self.subTest(evidence=evidence):
                source = self.questions / filename
                source.write_text("امتحان ٢٠٢٣\n1. Explain dehydration.\n", encoding="utf-8")
                response = self._written_response(source, year=2023, year_evidence=evidence)
                with patch("agy_exam_index.agy_writer.request_json", return_value=response):
                    index = agy_exam_index.build_index(self.module, [source])
                question = next(iter(index["questions"].values()))
                self.assertEqual(question["years"], [expected] if expected else [])
                rejected = evidence in {"2019", "unprinted heading 2023"} or filename.startswith("CamScanner")
                self.assertEqual(question["needs_review"], rejected)
                self.assertEqual(bool(question["occurrences"][0]["review_reason"]), rejected)
                if not rejected:
                    self.assertEqual(question["occurrences"][0]["year_source"], "paper" if filename.startswith("CamScanner") else "question")
                (self.questions / agy_exam_index.STATE_NAME).unlink()

    def test_header_year_outside_question_units_overrides_scan_filename(self):
        source = self.questions / "CamScanner ٢٩-١١-٢٠٢٣ ١٤.٤٧.txt"
        source.write_text("Date of exam: Tuesday:          1 8/١/٢٠٢٢\n1. Explain dehydration.\n", encoding="utf-8")
        units = agy_exam_index.read_source_units(source)
        response = self._written_response(
            source, unit_ids=[units[1].id], year=2022,
            year_evidence="Date of exam: Tuesday: 1 8/1/2022",
        )
        with patch("agy_exam_index.agy_writer.request_json", return_value=response):
            index = agy_exam_index.build_index(self.module, [source])
        with patch("agy_exam_index.agy_writer.request_json", side_effect=AssertionError("Use cached payload")):
            self.assertEqual(agy_exam_index.build_index(self.module, [source]), index)
        question = next(iter(index["questions"].values()))
        self.assertEqual(question["years"], [2022])
        self.assertEqual(question["occurrences"][0]["year_source"], "paper")
        self.assertFalse(question["needs_review"])
        self.assertEqual(index["sources"][0]["years"], [2022])
        self.assertEqual(index["sources"][0]["kind"], "past_exam")
        self.assertEqual(index["sources"][0]["needs_review"], 0)
        self.assertEqual(index["sources"][0]["review_reasons"], [])
        write_index(index, self.questions)
        with patch("remote_inventory.module_inventory", return_value=SimpleNamespace(
            sources=[], warning=None, remote_as_of=None, available=False,
        )):
            inventory = list_module_files(self.module)["files"]
        self.assertEqual(inventory[0]["question_review_count"], 0)
        self.assertEqual(inventory[0]["review_reasons"], index["sources"][0]["review_reasons"])


    def test_header_year_evidence_from_another_batch_is_verified(self):
        source = self.questions / "paper.txt"
        source.write_text("Exam 2022\n" + "Source note.\n" * 6 + "1. Explain dehydration.\n", encoding="utf-8")
        question = self._written_response(source, unit_ids=["L00008"], year=2022, year_evidence="Exam 2022")["questions"][0]

        def response(prompt, _schema, **_options):
            units = json.loads(prompt.split("Evidence units: ", 1)[1])
            core_ids = [unit["id"] for unit in units if unit["coverage"] == "required"]
            return {"covered_unit_ids": core_ids, "questions": [question] if "L00008" in core_ids else []}

        with patch("agy_exam_index.index_settings", return_value=(150, 1)), patch("agy_exam_index.agy_writer.request_json", side_effect=response):
            index = agy_exam_index.build_index(self.module, [source])
        extracted = next(iter(index["questions"].values()))
        self.assertEqual(extracted["years"], [2022])
        self.assertEqual(extracted["occurrences"][0]["year_source"], "paper")
        self.assertFalse(extracted["needs_review"])

    def test_paper_year_on_a_later_page_survives_provenance_only_when_printed_earlier(self):
        source = self.questions / "paper.txt"
        source.write_text("Final exam 2022\n--- Page 2 ---\n1. Explain dehydration.\n", encoding="utf-8")
        units = agy_exam_index.read_source_units(source)
        response = self._written_response(source, unit_ids=[units[-1].id], year=None)
        with patch("agy_exam_index.agy_writer.request_json", return_value=response):
            index = agy_exam_index.build_index(self.module, [source])
        occurrence = next(iter(index["questions"].values()))["occurrences"][0]
        self.assertEqual((occurrence["year"], occurrence["year_source"]), (2022, "paper"))
        verified = paper_backed_index(index, {}, self.module.module_id, self.questions)
        self.assertEqual(next(iter(verified["questions"].values()))["years"], [2022])
        occurrence["year"] = 2019
        self.assertEqual(paper_backed_index(index, {}, self.module.module_id, self.questions)["questions"], {})

    def test_single_paper_year_fallback_preserves_question_review_semantics(self):
        cases = [
            ("End 2025.txt", "", None, None, 2025, "file_name", False),
            ("scan 2023.txt", "Date of exam: 18/1/2022\n", None, None, 2023, "file_name", False),
            ("paper.txt", "Exam 2022\n", 2020, "Unprinted exam 2020", 2022, "paper", True),
            ("paper.txt", "", 2020, "Unprinted exam 2020", None, None, True),
            ("End 2025.txt", "Exams 2022 and 2023\n", None, None, 2025, "file_name", False),
            ("End 2022 2025.txt", "Exams 2022\n", None, None, None, None, False),
            ("Bank 2025.txt", "", None, None, None, None, False),
            ("End 2024 2025.txt", "", None, None, None, None, False),
        ]
        for filename, header, year, evidence, expected, provenance, review in cases:
            with self.subTest(filename=filename, header=header, year=year):
                source = self.questions / filename
                source.write_text(header + "1. Explain dehydration.\n", encoding="utf-8")
                response = self._written_response(source, year=year, year_evidence=evidence)
                with patch("agy_exam_index.agy_writer.request_json", return_value=response):
                    index = agy_exam_index.build_index(self.module, [source])
                question = next(iter(index["questions"].values()))
                self.assertEqual(question["years"], [expected] if expected else [])
                occurrence = question["occurrences"][0]
                self.assertEqual(occurrence["year"], expected)
                self.assertEqual(occurrence["year_source"], provenance)
                self.assertEqual(question["needs_review"], review)
                self.assertEqual(bool(occurrence["review_reason"]), review)
                if review:
                    self.assertIn("cannot be verified", occurrence["review_reason"])
                (self.questions / agy_exam_index.STATE_NAME).unlink()

    def test_question_bank_keeps_only_verified_question_years(self):
        source = self.questions / "bank.txt"
        source.write_text("Exam 2022\n1. Explain dehydration.\nExam 2023\n2. Describe fluid loss.\n", encoding="utf-8")
        units = agy_exam_index.read_source_units(source)
        response = self._written_response(source, unit_ids=[units[1].id], year=2022, year_evidence="Exam 2022")
        response["questions"].append({
            **response["questions"][0], "unit_ids": [units[3].id], "stem": "Describe fluid loss.",
            "number": 2, "year": None, "year_evidence": None,
        })
        with patch("agy_exam_index.agy_writer.request_json", return_value=response):
            index = agy_exam_index.build_index(self.module, [source])
        questions = {question["number"]: question for question in index["questions"].values()}
        self.assertEqual(questions[1]["years"], [2022])
        self.assertEqual(questions[2]["years"], [])
        self.assertIsNone(questions[2]["occurrences"][0]["year_source"])
        self.assertFalse(questions[2]["needs_review"])
        self.assertEqual(index["sources"][0]["kind"], "question_bank")
        self.assertEqual(index["sources"][0]["years"], [2022, 2023])

    def test_ungrounded_stem_is_dropped_and_counted_while_valid_questions_resume(self):
        source = self.questions / "paper.txt"
        source.write_text("1. Explain dehydration.\n", encoding="utf-8")
        response = self._written_response(source)
        response["questions"].append({**response["questions"][0], "stem": "Describe unrelated quantum physics."})
        with patch("agy_exam_index.agy_writer.request_json", return_value=response):
            index = agy_exam_index.build_index(self.module, [source])
        with patch("agy_exam_index.agy_writer.request_json", side_effect=AssertionError("Must resume cached batch")):
            resumed = agy_exam_index.build_index(self.module, [source])
        self.assertEqual(index, resumed)
        self.assertEqual([q["stem"] for q in index["questions"].values()], ["Explain dehydration."])
        self.assertEqual(index["sources"][0]["dropped_questions"], 1)
        self.assertEqual(index["sources"][0]["dropped_question_reasons"], {"A question stem is not grounded": 1})

    def test_invalid_options_and_metadata_do_not_abort_other_questions(self):
        source = self.questions / "paper.txt"
        source.write_text("1. Explain dehydration.\nA. Fluid loss\nB. Reduced intake\n", encoding="utf-8")
        response = self._written_response(
            source, kind="mcq", number="bad", section=[], topic=42, review_reason={}, needs_review="false",
            correct_option="c", answer_unit_ids=["missing"],
            options=[{"label": "a", "text": "Fluid loss"}, {"label": "b", "text": "Reduced intake"},
                     {"label": "c", "text": "Quantum tunnelling"}],
        )
        response["questions"].append({**response["questions"][0], "options": [{"label": "a", "text": "Fluid loss"}]})
        with patch("agy_exam_index.agy_writer.request_json", return_value=response):
            index = agy_exam_index.build_index(self.module, [source])
        question = next(iter(index["questions"].values()))
        self.assertEqual(question["options"], {"a": "Fluid loss", "b": "Reduced intake"})
        self.assertIsNone(question["number"])
        self.assertIsNone(question["topic"])
        self.assertIsNone(question["answer"])
        self.assertEqual(question["answer_status"], "unverified")
        self.assertTrue(question["needs_review"])
        self.assertIn("answer cannot be verified", question["occurrences"][0]["review_reason"])
        self.assertEqual(index["sources"][0]["dropped_questions"], 1)

    def test_structurally_failed_file_does_not_prevent_another_file_being_indexed(self):
        first, second = self.questions / "A.txt", self.questions / "B.txt"
        for source in (first, second):
            source.write_text("1. Explain dehydration.\n", encoding="utf-8")
        malformed = agy_exam_index.agy_writer.AgyProposalError("invalid JSON", "{broken")
        def response(prompt, _schema, **_options):
            if f"Source file: {first.name}\n" in prompt:
                raise malformed
            return self._written_response(second)

        with patch("agy_exam_index.agy_writer.request_json", side_effect=response):
            index = agy_exam_index.build_index(self.module, [first, second])
        self.assertEqual([source["file"] for source in index["sources"]], ["B.txt"])
        failure = index["failed_sources"][0]
        self.assertEqual(failure["file"], "A.txt")
        self.assertEqual(failure["sha256"], hashlib.sha256(first.read_bytes()).hexdigest())
        self.assertIn("after one repair attempt", failure["error"])
        with patch("agy_exam_index.agy_writer.request_json", return_value=self._written_response(first)) as request:
            retried = agy_exam_index.build_index(self.module, [first, second])
        self.assertEqual(request.call_count, 1)
        self.assertEqual(len(retried["sources"]), 2)
        self.assertEqual(retried["failed_sources"], [])

    def test_later_batch_failure_keeps_earlier_validated_batches_for_retry(self):
        source = self.questions / "paper.txt"
        source.write_text("\n".join(f"Source note {number}: " + "evidence " * 25 for number in range(3)), encoding="utf-8")
        seen = []

        def response(prompt, _schema, **_options):
            units = json.loads(prompt.split("Evidence units: ", 1)[1])
            core_ids = [unit["id"] for unit in units if unit["coverage"] == "required"]
            seen.append(core_ids)
            return {"covered_unit_ids": core_ids, "questions": []}

        def fail_second(prompt, schema, **options):
            if seen:
                return {"covered_unit_ids": [], "questions": []}
            return response(prompt, schema, **options)

        with patch("agy_exam_index.index_settings", return_value=(260, 1)):
            with patch("agy_exam_index.agy_writer.request_json", side_effect=fail_second):
                failed = agy_exam_index.build_index(self.module, [source])
            self.assertEqual(failed["sources"], [])
            self.assertEqual(failed["failed_sources"][0]["file"], source.name)
            state = json.loads((self.questions / agy_exam_index.STATE_NAME).read_text())
            self.assertEqual(list(state["files"][source.name]["batches"]), ["0"])
            with patch("agy_exam_index.agy_writer.request_json", side_effect=response):
                retried = agy_exam_index.build_index(self.module, [source])
        self.assertEqual(seen, [["L00001"], ["L00002"], ["L00003"]])
        self.assertEqual(retried["failed_sources"], [])
        self.assertEqual(retried["sources"][0]["file"], source.name)

    def test_incomplete_model_coverage_records_file_failure(self):
        source = self.questions / "questions.txt"
        source.write_text("One source line.\nSecond source line.\n", encoding="utf-8")
        units = agy_exam_index.read_source_units(source)
        payload = {"covered_unit_ids": [units[0].id], "questions": []}
        with patch("agy_exam_index.agy_writer.request_json", return_value=payload):
            index = agy_exam_index.build_index(self.module, [source])
        self.assertEqual(index["sources"], [])
        self.assertIn("cover every source unit", index["failed_sources"][0]["error"])

    def test_repeated_or_question_cited_coverage_needs_no_retry(self):
        source = self.questions / "questions.txt"
        source.write_text("One source line.\nSecond source line.\n", encoding="utf-8")
        units = agy_exam_index.read_source_units(source)
        payload = {"covered_unit_ids": [units[0].id, units[0].id], "questions": [{"unit_ids": [units[1].id]}]}
        with patch("agy_exam_index.agy_writer.request_json", return_value=payload) as request:
            index = agy_exam_index.build_index(self.module, [source])
        self.assertEqual(index["failed_sources"], [])
        self.assertEqual(request.call_count, 1)

    def test_skipped_units_return_to_agy_in_smaller_groups(self):
        source = self.questions / "paper.txt"
        source.write_text("\n".join(f"Source note {number}: " + "evidence " * 25 for number in range(3)), encoding="utf-8")
        seen = []

        def response(prompt, _schema, **_options):
            units = json.loads(prompt.split("Evidence units: ", 1)[1])
            core_ids = [unit["id"] for unit in units if unit["coverage"] == "required"]
            seen.append(core_ids)
            # The first whole-paper reading skips its last unit.
            return {"covered_unit_ids": core_ids[:-1] if len(seen) == 1 else core_ids, "questions": []}

        with patch("agy_exam_index.index_settings", return_value=(100_000, 1)):
            with patch("agy_exam_index.agy_writer.request_json", side_effect=response):
                index = agy_exam_index.build_index(self.module, [source])
        self.assertEqual(index["failed_sources"], [])
        self.assertEqual(index["sources"][0]["file"], source.name)
        self.assertEqual(len(seen[0]), 3)
        self.assertGreater(len(seen), 1)
        self.assertTrue(all(len(core) < 3 for core in seen[1:]))

    def test_build_exam_index_cli_keeps_previous_publication_when_all_files_fail(self):
        import run_transcription

        config = Path(self.temporary.name) / "config.json"
        config.write_text("{}", encoding="utf-8")
        first, second = self.questions / "A.txt", self.questions / "B.txt"
        for source in (first, second):
            source.write_text("1. Explain dehydration.\n", encoding="utf-8")
        broken = {"covered_unit_ids": [], "questions": []}
        argv = ["run_transcription.py", "--workspace", self.temporary.name, "--module", "pediatric", "--build-exam-index"]
        for all_failed in (False, True):
            with self.subTest(all_failed=all_failed):
                (self.questions / agy_exam_index.STATE_NAME).unlink(missing_ok=True)
                target = self.questions / "exam-index.json"
                previous_bytes = target.read_bytes() if target.exists() else None
                stdout, stderr = io.StringIO(), io.StringIO()
                second_response = broken if all_failed else self._written_response(second)
                with patch.dict(os.environ, {"TRANSCRIBER_CONFIG_PATH": str(config), "TRANSCRIBER_WORKSPACE": self.temporary.name}), patch.object(
                    sys, "argv", argv,
                ), patch("agy_exam_index.agy_writer.request_json", side_effect=lambda prompt, _schema, answer=second_response, **_options: broken if "Source file: A.txt\n" in prompt else answer), redirect_stdout(stdout), redirect_stderr(stderr):
                    exit_code = run_transcription.main()
                self.assertEqual(exit_code, 1 if all_failed else 0, stderr.getvalue())
                self.assertIn("Failed: A.txt", stdout.getvalue())
                state = json.loads((self.questions / agy_exam_index.STATE_NAME).read_text())
                self.assertEqual(len(state["failed_sources"]), 2 if all_failed else 1)
                if all_failed:
                    self.assertIn("Failed: B.txt", stdout.getvalue())
                    self.assertEqual(target.read_bytes(), previous_bytes)
                else:
                    published = json.loads(target.read_text())
                    self.assertEqual(len(published["sources"]), 1)
                    self.assertEqual(len(published["failed_sources"]), 1)

    def test_build_exam_index_cli_without_papers_keeps_previous_publication(self):
        import run_transcription

        config = Path(self.temporary.name) / "config.json"
        config.write_text("{}", encoding="utf-8")
        target = self.questions / "exam-index.json"
        previous_bytes = b'{"schema_version": 2, "sources": [], "questions": {}}\n'
        target.write_bytes(previous_bytes)
        argv = ["run_transcription.py", "--workspace", self.temporary.name, "--module", "pediatric", "--build-exam-index"]
        stdout, stderr = io.StringIO(), io.StringIO()
        with patch.dict(os.environ, {"TRANSCRIBER_CONFIG_PATH": str(config), "TRANSCRIBER_WORKSPACE": self.temporary.name}), patch.object(
            sys, "argv", argv,
        ), patch("agy_exam_index.agy_writer.request_json", side_effect=AssertionError("No papers to extract")), redirect_stdout(stdout), redirect_stderr(stderr):
            exit_code = run_transcription.main()
        self.assertEqual(exit_code, 1)
        self.assertIn("No extractable question files", stderr.getvalue())
        self.assertEqual(target.read_bytes(), previous_bytes)

    def test_retrying_selected_papers_keeps_other_published_failures(self):
        import run_transcription

        first, second, third = (self.questions / name for name in ("A.txt", "B.txt", "C.txt"))
        for source in (first, second, third):
            source.write_text(f"Source note from {source.stem}.\n", encoding="utf-8")
        requested = []
        def covered(prompt, _schema, **_options):
            requested.append(next(path.name for path in (first, second, third) if f"Source file: {path.name}\n" in prompt))
            if requested[-1] in failing:
                raise agy_exam_index.agy_writer.AgyWriterError("agy failed: unsupported request")
            units = json.loads(prompt.split("Evidence units: ", 1)[1])
            return {"covered_unit_ids": [unit["id"] for unit in units if unit["coverage"] == "required"], "questions": []}
        failing = {"A.txt", "B.txt"}
        with patch("agy_exam_index.agy_writer.request_json", side_effect=covered):
            run_transcription.build_exam_index(self.module, {"exam_index_workers": 1})
            before = json.loads((self.questions / "exam-index.json").read_text())["failed_sources"]
            failing, requested[:] = set(), []
            index, _target = run_transcription.build_exam_index(self.module, {"exam_index_workers": 1}, papers=["A.txt"])
        self.assertEqual(requested, ["A.txt"])
        self.assertEqual([source["file"] for source in index["sources"]], ["A.txt", "C.txt"])
        self.assertEqual(index["failed_sources"], [failure for failure in before if failure["file"] == "B.txt"])
        second.write_text("Source note from B, corrected.\n", encoding="utf-8")
        requested[:] = []
        with patch("agy_exam_index.agy_writer.request_json", side_effect=covered):
            index, _target = run_transcription.build_exam_index(self.module, {"exam_index_workers": 1}, papers=[])
        # A changed paper is read again even when no failed paper was chosen.
        self.assertEqual(requested, ["B.txt"])
        self.assertEqual(index["failed_sources"], [])

    def test_preparation_failure_is_published_while_readable_paper_is_indexed(self):
        import run_transcription

        first, second = self.questions / "A.txt", self.questions / "B.txt"
        first.write_bytes(b"\xff\xfe")
        second.write_text("1. Explain dehydration.\n", encoding="utf-8")
        ready, failures = run_transcription._prepare_exam_sources(self.module, {})
        self.assertEqual(ready, [second])
        self.assertEqual(failures[0]["file"], "A.txt")
        with patch("agy_exam_index.agy_writer.request_json", return_value=self._written_response(second)):
            index = agy_exam_index.build_index(self.module, ready, failed_sources=failures)
            published, target = run_transcription._publish_exam_index(self.module, ready, index)
        self.assertEqual(published["failed_sources"], failures)
        self.assertEqual(json.loads(target.read_text())["failed_sources"], failures)

    def test_manual_repairs_do_not_retain_occurrences_after_prepared_text_changes(self):
        source = self.questions / "Exam.pdf"
        source.write_bytes(b"unchanged original")
        original_hash = hashlib.sha256(b"unchanged original").hexdigest()
        old = {
            "schema_version": 2, "extractor": "agy", "module": self.module.module_id,
            "sources": [{"file": source.name, "sha256": original_hash, "units_sha256": "old-units"}],
            "questions": {"repaired": {
                "stem": "Reviewed question", "repaired_by_hand": True,
                "occurrences": [{"source": source.name, "unit_ids": ["P00001"]}],
            }},
        }
        write_index(old, self.questions)
        fresh = {
            "schema_version": 2, "extractor": "agy", "module": self.module.module_id,
            "sources": [{"file": source.name, "sha256": original_hash, "units_sha256": "new-units"}],
            "questions": {},
        }

        rebuilt = carry_over_repairs(fresh, self.questions)

        self.assertNotIn("repaired", rebuilt["questions"])


if __name__ == "__main__":
    unittest.main()
